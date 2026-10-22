"""Export saved benchmark scores; never regenerate SQL or query the database."""
from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE


def truth(value):
    return str(value).strip().lower() in {'true', '1', '1.0'}


def export_answers(frame, run, out):
    run, out = Path(run), Path(out)
    out.mkdir(parents=True, exist_ok=True)
    gt = pd.read_excel(run / 'inputs/ground_truth.xlsx', keep_default_na=False)
    reference = dict(zip(gt.iloc[:, 1].astype(str).str.strip(), gt.iloc[:, 2]))
    records = []
    for row in frame.sort_values('question_id').to_dict('records'):
        available = truth(row['evaluation_available'])
        matched = truth(row['result_values_match'])
        records.append({
            '问题 (Query)': row['question'],
            '预测的SQL (Predicted)': row['final_sql'],
            '正确的SQL (Ground Truth)': row.get('gt_sql') or reference.get(str(row['question']).strip(), ''),
            '是否正确 (Is_Correct)': matched if available else None,
            '预测执行错误 (Predicted Error)': row.get('error', ''),
            '标准SQL执行错误 (GT Execution Error)': row.get('gt_error', ''),
            '预测的返回结果 (Predicted JSON)': row.get('predicted_result', '[]'),
            '正确的返回结果 (Correct JSON)': row.get('gt_result', '[]'),
            '题号 (Question_ID)': int(row['question_id']),
            '评测可用 (Evaluation_Available)': available,
            '列名—值对应匹配 (Named_Match)': truth(row['strict_match']) if available else None,
            '评测状态': ('正确' if matched else '错误') if available else '无法评测',
            '进程状态': row['process_status'],
        })
    result = pd.DataFrame(records)
    csv_path, xlsx_path = out / '问答结果汇总.csv', out / '问答结果汇总.xlsx'
    result.to_csv(csv_path, index=False, encoding='utf-8-sig')
    book = Workbook()
    sheet = book.active
    sheet.title = '问答结果'
    sheet.append(list(result.columns))
    clipped = 0
    for record in records:
        values = []
        for value in record.values():
            if isinstance(value, str):
                value = ILLEGAL_CHARACTERS_RE.sub('', value)
                if len(value) > 32767:
                    clipped += 1
                    value = value[:32000] + '\n[Excel单元格长度受限，完整内容见同目录问答结果汇总.csv]'
            values.append(value)
        sheet.append(values)
        # Preserve SQL, questions and JSON as text even if they start with '='.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    sheet.freeze_panes = 'D2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='245A81')
    for column, width in {'A': 48, 'B': 65, 'C': 65, 'D': 18, 'E': 28, 'F': 28,
                          'G': 65, 'H': 65, 'I': 12, 'J': 18, 'K': 28, 'L': 14, 'M': 18}.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 48
    notes = book.create_sheet('口径说明')
    notes.append(['项目', '说明'])
    for pair in [
        ('是否正确', '采用89%报告的结果值匹配：忽略别名和行序，保留列序、精确值、NULL和重复行。'),
        ('列名—值对应匹配', '单独保留原逐题evaluation.xlsx的Is_Correct口径，忽略行序和列序；与本表是否正确不同。'),
        ('无法评测', '无标准答案或标准SQL执行失败时，是否正确留空，评测状态为无法评测；统计准确率时仍占全部题目的分母。'),
        ('答案', '标准答案和生成答案均提供SQL及已记录的返回结果。导出不重新调用模型或数据库。'),
        ('Excel长文本', f'{clipped}个单元格超出Excel限制，已标注；CSV保留完整原始文本。'),
        ('原始记录', '逐题原始evaluation.json和evaluation.xlsx保留，不修改既有评分。'),
    ]:
        notes.append(pair)
    notes.column_dimensions['A'].width = 24
    notes.column_dimensions['B'].width = 110
    book.save(xlsx_path)
    return csv_path, xlsx_path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('run', type=Path)
    parser.add_argument('--output', type=Path, help='Defaults to RUN/analysis')
    args = parser.parse_args()
    out = args.output or args.run / 'analysis'
    frame = pd.read_csv(args.run / 'analysis/逐题结果.csv', keep_default_na=False)
    for path in export_answers(frame, args.run, out):
        print(path)


if __name__ == '__main__':
    main()
